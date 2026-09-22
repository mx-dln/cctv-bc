from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "CCTV_BLOCKCHAIN_USER_MANUAL.docx"


def set_run(run, size=10.5, bold=False, color="000000"):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        set_run(run, size=15 if level == 1 else 12.5, bold=True)
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(4)
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    set_run(run)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    set_run(run)
    return paragraph


def add_numbered(doc, number, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.first_line_indent = Inches(-0.22)
    run = paragraph.add_run(f"{number}.  {text}")
    set_run(run)
    return paragraph


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_table(doc, rows, widths):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0].cells
    for i, value in enumerate(rows[0]):
        header[i].text = value
        header[i].width = Inches(widths[i])
        header[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(header[i], "303030")
        for paragraph in header[i].paragraphs:
            for run in paragraph.runs:
                set_run(run, bold=True, color="FFFFFF")

    for row in rows[1:]:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].width = Inches(widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if len(table.rows) % 2 == 1:
                shade_cell(cells[i], "F6F7F9")
            for paragraph in cells[i].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run(run, size=9.5)
    doc.add_paragraph()
    return table


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("CCTV Blockchain Chain of Custody User Manual")
    set_run(title_run, size=18, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run("FICOBank capstone prototype using private footage storage, SHA 256 hashing, and Hyperledger Fabric")
    set_run(sub_run, size=10.5, color="444444")

    add_heading(doc, "System Flow", 1)
    add_body(doc, "The implemented workflow is Authorized Footage Source, then Local or Private Storage, then SHA 256 Hash, then Hyperledger Fabric, then Verification. The authorized source may be an uploaded CCTV file, exported DVR/NVR footage, USB camera capture, or controlled footage.")
    add_body(doc, "The system is not a replacement for CCTV cameras, DVR, or NVR, and it does not strictly require a DVR/NVR for prototype testing. It is a blockchain augmented chain of custody layer for authorized footage. The footage file stays in local or private storage. Hyperledger Fabric stores the custody proof, including the record ID, footage hash, metadata hash, final SHA 256 hash, timestamp, source or camera reference, transaction ID, and accountable user details.")

    add_table(
        doc,
        [
            ["Stage", "Purpose"],
            ["Authorized Footage Source", "Provides the footage. This can be uploaded CCTV footage, exported DVR/NVR footage, USB camera capture, or controlled footage."],
            ["Optional DVR or NVR", "One possible production source, but not required for prototype registration or verification."],
            ["Local or Private Storage", "Stores the actual video file outside the blockchain."],
            ["SHA 256 Hash", "Creates a digital fingerprint of the footage and custody metadata."],
            ["Hyperledger Fabric", "Stores the immutable custody proof for future verification."],
            ["Verification", "Recomputes the current hash and compares it with the original Fabric proof."],
        ],
        [1.7, 5.6],
    )

    add_heading(doc, "Login", 1)
    add_body(doc, "Use the administrator account for testing and demonstration.")
    add_bullet(doc, "Email: admin@ficobank.com")
    add_bullet(doc, "Password: password")

    add_heading(doc, "Dashboard", 1)
    add_body(doc, "The dashboard is the main control panel. It shows the evidence count, Fabric commits, pending verification count, alerts, the evidence flow, recent evidence, Fabric proofs, and recent audit activity.")
    add_bullet(doc, "Click Register Evidence to add new CCTV footage to the chain of custody.")
    add_bullet(doc, "Click Verify Evidence to check whether a file still matches its recorded proof.")

    add_heading(doc, "Cameras", 1)
    add_body(doc, "Use the Cameras page only when checking a connected DVR/NVR or mock camera provider. A DVR/NVR is not required for prototype evidence registration because uploaded CCTV footage, USB camera capture, and controlled footage are supported.")
    add_body(doc, "If the camera list is empty, run this command from the Laravel project folder:")
    add_body(doc, "php artisan cctv:discover-cameras")

    add_heading(doc, "Evidence Registration", 1)
    for index, step in enumerate([
        "Open Evidence Register.",
        "Select the camera that produced the footage.",
        "Upload a CCTV video file.",
        "Fill in filename, resolution, date and time, duration, and recording information.",
        "Leave Already stored footage path blank when uploading a new file.",
        "Click Register, Hash, and Commit.",
        "Wait until the record displays Fabric committed.",
    ], start=1):
        add_numbered(doc, index, step)

    add_heading(doc, "What Happens During Registration", 1)
    add_bullet(doc, "Laravel saves the uploaded footage in private local storage.")
    add_bullet(doc, "The system computes a SHA 256 hash for the footage.")
    add_bullet(doc, "The system computes a SHA 256 hash for the custody metadata.")
    add_bullet(doc, "The final custody proof is committed to Hyperledger Fabric.")
    add_bullet(doc, "The database stores the custody record and Fabric transaction ID.")

    add_heading(doc, "Evidence Verification", 1)
    for index, step in enumerate([
        "Open Evidence Verification.",
        "Choose the evidence record to verify.",
        "Click Verify Stored Original to check the private stored footage.",
        "To test possible tampering, upload a questioned, cut, or edited file as the comparison file.",
        "Click Compare.",
        "Review the local and blockchain verification result.",
    ], start=1):
        add_numbered(doc, index, step)

    add_body(doc, "Important: do not register an edited or cut video as new evidence when testing tampering. If it is registered as new evidence, it receives its own new hash and can verify as its own original record. For tamper testing, register the original footage once, then upload the edited or cut video only in Evidence Verification as a comparison file.")

    add_table(
        doc,
        [
            ["Result", "Meaning"],
            ["Verified", "The footage and metadata match the Fabric custody proof."],
            ["Tampered", "The current hash does not match the recorded hash."],
            ["Missing", "The local or private footage file is unavailable."],
            ["Unavailable", "The Fabric gateway or ledger proof cannot be reached."],
        ],
        [1.5, 5.8],
    )

    add_heading(doc, "Fabric Ledger", 1)
    add_body(doc, "Use the Fabric Ledger page to view blockchain registration and verification records. Registration rows show the Fabric commit proof. Verification rows show whether the stored original or uploaded comparison matched the blockchain recorded proof. A real Fabric transaction ID is a long hexadecimal value. Older TX style IDs are from earlier demo or simulation records.")

    add_heading(doc, "Alerts", 1)
    add_body(doc, "Use the Alerts page to review discrepancy or integrity warnings, including possible tampering, failed verification, missing footage, or failed blockchain commits.")

    add_heading(doc, "Audit Reports", 1)
    add_body(doc, "Use Audit Reports to export chain of custody reports for documentation, review, or capstone demonstration.")

    add_heading(doc, "Settings", 1)
    add_body(doc, "Use Settings to configure optional CCTV providers, activate Mock Provider for demonstration, configure a Dahua DVR or NVR for production style testing if available, and manage user and security settings. When no DVR/NVR is available, use controlled footage upload or USB camera capture.")

    add_heading(doc, "Recommended Demonstration Workflow", 1)
    demo_steps = [
        "Start Docker Desktop.",
        "Start the Hyperledger Fabric network in Ubuntu WSL.",
        "Start the Node Fabric gateway.",
        "Open the Laravel application at https://cctv-bc.test.",
        "Log in using the administrator account.",
        "Open Settings and activate Mock Provider if needed.",
        "Run php artisan cctv:discover-cameras if the camera list is empty.",
        "Open Evidence Register and upload a sample MP4.",
        "Click Register, Hash, and Commit.",
        "Confirm the record displays Fabric committed.",
        "Open Evidence Verification and click Verify Stored Original.",
        "Upload the same video with a different filename as a comparison file. It should verify because the video content hash is the same.",
        "Upload a cut or edited version as a comparison file. It should show tampered because the content hash changed.",
        "Open Fabric Ledger and Audit Reports for demonstration evidence.",
    ]
    for index, step in enumerate(demo_steps, start=1):
        add_numbered(doc, index, step)

    add_heading(doc, "Defense Guide", 1)
    add_body(doc, "Short explanation: Ang system namin ay hindi replacement ng CCTV, DVR, or NVR, and hindi rin required ang DVR/NVR for prototype testing. It is a blockchain augmented chain of custody layer for authorized footage. The footage remains in local or private storage, while the system generates SHA 256 hashes and records the integrity proof in Hyperledger Fabric. During verification, the system recomputes the hash and compares it with the blockchain recorded proof to detect alteration, deletion, or mismatch.")

    add_heading(doc, "System Flow To Defend", 2)
    add_body(doc, "Authorized footage source, such as uploaded CCTV file, exported DVR/NVR footage, USB camera capture, or controlled footage -> Local or Private Storage -> SHA 256 Hash -> Hyperledger Fabric -> Verification")
    for index, step in enumerate([
        "Authorized footage is provided from an uploaded CCTV file, exported DVR/NVR footage, USB camera capture, or controlled source.",
        "Footage is stored locally or privately.",
        "Authorized user registers the footage in the system.",
        "The system generates a SHA 256 hash.",
        "The hash and metadata proof are committed to Hyperledger Fabric.",
        "During verification, the system hashes the current file again.",
        "If hashes match, the evidence is verified.",
        "If hashes do not match, the system flags possible tampering.",
    ], start=1):
        add_numbered(doc, index, step)

    add_heading(doc, "Panel Questions", 2)
    qa_rows = [
        ["Question", "Defense answer"],
        ["Bakit hindi video mismo sa blockchain?", "Hindi practical mag-store ng video sa blockchain because CCTV files are large. Blockchain is best used as an immutable proof layer, not as video storage. The correct design is to store the video securely in private storage and store its SHA 256 hash and metadata proof in Hyperledger Fabric."],
        ["Need ba talaga DVR/NVR?", "No. DVR/NVR is not strictly required for the prototype. The system needs authorized footage bytes that can be stored privately, hashed, and verified. A DVR/NVR is one possible source for production, but uploaded CCTV footage, exported footage, USB camera capture, or controlled footage can also be used for capstone testing."],
        ["Mock DVR lang ba yan?", "The Mock Provider is optional and only for controlled demonstration and testing. It simulates camera channels so the workflow can be tested without exposing real bank CCTV infrastructure. The system can also use uploaded footage or USB camera capture without a DVR/NVR."],
        ["Ano role ng Hyperledger Fabric?", "Hyperledger Fabric stores the immutable custody proof: evidence hash, metadata hash, record ID, timestamp, camera reference, and transaction ID."],
        ["Paano nadedetect tampering?", "Through SHA 256 hash comparison. If even one small part of the video changes, the generated hash changes. If the current hash does not match the Fabric proof, the evidence is flagged as altered or tampered."],
        ["What if unavailable ang verification?", "The system separates local and blockchain verification. This tells the user whether the problem is storage side, missing footage, or Fabric gateway availability."],
        ["Is this 100 percent deployed sa FICOBank?", "This is a capstone prototype aligned with the proposed architecture. Actual production deployment may integrate with FICOBank's real DVR/NVR or another approved footage export/storage workflow, depending on the bank's infrastructure."],
    ]
    add_table(doc, qa_rows, [2.15, 5.15])

    add_heading(doc, "Defense Demo Script", 2)
    demo_script = [
        "Open Dashboard and say: This shows the chain of custody control panel and the implemented evidence flow.",
        "Open Evidence Register and say: Here, authorized CCTV footage can be uploaded or captured from a controlled camera source.",
        "Upload or record a short video and say: The footage is stored privately. The system will not upload the video to blockchain.",
        "Click Register, Hash, and Commit and say: At this point, the system generates SHA 256 hashes and commits the proof to Hyperledger Fabric.",
        "Show Fabric committed and say: This confirms the blockchain proof was recorded.",
        "Open Fabric Ledger and say: This transaction ID is the Fabric proof of registration.",
        "Open Evidence Verification and click Verify Stored Original. Say: The system recomputes the hash and compares it with the stored blockchain proof.",
        "Show verified result and say: A verified result means the current footage matches the original registered proof.",
        "Upload a modified video as a comparison file and say: This altered file produces a different SHA 256 hash, so the system detects a mismatch. The system identifies evidence through custody record ID and blockchain recorded proof, not filename alone.",
    ]
    for index, step in enumerate(demo_script, start=1):
        add_numbered(doc, index, step)

    add_heading(doc, "Closing Answer", 2)
    add_body(doc, "Our system strengthens CCTV evidence integrity by adding a tamper evident blockchain proof layer. It does not replace the bank's CCTV infrastructure. Instead, it secures the chain of custody after authorized footage is exported or registered. The footage stays in private storage, while SHA 256 hashes and metadata proofs are stored in Hyperledger Fabric. This makes evidence verification traceable, auditable, and resistant to undetected alteration.")

    add_heading(doc, "Troubleshooting", 1)
    add_table(
        doc,
        [
            ["Issue", "Fix"],
            ["Fabric pending", "Check Docker Desktop, WSL Fabric containers, Node gateway, and curl.exe http://127.0.0.1:8787/health."],
            ["Verification unavailable", "Check gateway health, BLOCKCHAIN_SIMULATE=false, and whether the record was committed to Fabric."],
            ["Cameras empty", "Activate Mock Provider and run php artisan cctv:discover-cameras."],
            ["Upload failed", "Use a smaller MP4, restart PHP or Herd, and check PHP upload limits."],
        ],
        [2.0, 5.3],
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
